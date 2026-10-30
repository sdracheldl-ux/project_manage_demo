# -*- coding: utf-8 -*-
"""投促局阶段汇报 PPT 生成脚本（6页，16:9）"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

BASE = "/Users/kangsuzhen/Desktop/东湖项目/项目全周期管理"
ASSETS = os.path.join(BASE, "ppt_assets")
OUT = os.path.join(BASE, "ppt_assets", "投促局阶段汇报PPT.pptx")

# 配色
NAVY = RGBColor(0x16, 0x36, 0x5C)      # 主色深蓝
MID = RGBColor(0x2F, 0x55, 0x97)       # 次级蓝
GOLD = RGBColor(0xC6, 0x92, 0x14)      # 金色点缀
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)     # 浅灰背景
BLUE_BG = RGBColor(0xE8, 0xEE, 0xF7)   # 浅蓝背景
GRAY_BG = RGBColor(0xEF, 0xEF, 0xEF)   # 灰色背景（痛点）
DARK = RGBColor(0x33, 0x33, 0x33)      # 正文深灰
GRAY = RGBColor(0x59, 0x59, 0x59)      # 次级文字
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xD9, 0xE0, 0xEA)    # 图片边框
FONT = "PingFang SC"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         before=0, after=0, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return p


def rich(tf, parts, align=PP_ALIGN.LEFT, before=0, after=0, first=False):
    """parts: list of (text, size, bold, color)"""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    for text, size, bold, color in parts:
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return p


def card(s, x, y, w, h, fill=LIGHT, line=None, radius=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def title(s, text, sub=None):
    tf = tb(s, 0.5, 0.32, 12.33, 0.65)
    para(tf, text, 26, True, NAVY, first=True)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.52), Inches(1.02),
                            Inches(1.6), Inches(0.055))
    ln.fill.solid()
    ln.fill.fore_color.rgb = GOLD
    ln.line.fill.background()
    ln.shadow.inherit = False
    if sub:
        tf2 = tb(s, 2.3, 0.62, 10.5, 0.4)
        para(tf2, sub, 12, False, GRAY, first=True)


def pic(s, path, x, y, w=None, h=None):
    p = s.shapes.add_picture(os.path.join(ASSETS, path), Inches(x), Inches(y),
                             Inches(w) if w else None, Inches(h) if h else None)
    p.line.color.rgb = BORDER
    p.line.width = Pt(1)
    p.shadow.inherit = False
    return p


def page_no(s, n):
    tf = tb(s, 12.55, 7.08, 0.6, 0.35)
    para(tf, f"{n:02d}", 10, False, GRAY, align=PP_ALIGN.RIGHT, first=True)


# ============ P1 为什么建 ============
s = slide()
title(s, "为什么建：实地调研发现的真问题", "来源于对投促局的实地调研")

tf = tb(s, 0.5, 1.22, 5.0, 0.35)
para(tf, "调研发现 · 三大痛点", 13, True, GOLD, first=True)

pains = [
    ("早期项目无处管理",
     "大量谋划、在谈项目不具上报条件，市级系统无处管理——进度靠电话催、表格记，区级看不到全量"),
    ("考核统计依赖人工",
     "国企、部门、园区考核线下统计：认缴、见商出差、活动逐月人工汇总，月均耗费约 5 人天"),
    ("多头填报、口径不一",
     "同一项目市级报一次、区级表格再填一次，重复录入与统计口径不一并存"),
]
y = 1.62
for i, (t, d) in enumerate(pains, 1):
    card(s, 0.5, y, 7.1, 1.08, GRAY_BG)
    n = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.72), Inches(y + 0.33),
                           Inches(0.42), Inches(0.42))
    n.fill.solid(); n.fill.fore_color.rgb = NAVY; n.line.fill.background()
    n.shadow.inherit = False
    ntf = n.text_frame; ntf.word_wrap = False
    ntf.margin_left = ntf.margin_right = ntf.margin_top = ntf.margin_bottom = 0
    np_ = ntf.paragraphs[0]; np_.alignment = PP_ALIGN.CENTER
    nr = np_.add_run(); nr.text = str(i)
    nr.font.name = FONT; nr.font.size = Pt(16); nr.font.bold = True
    nr.font.color.rgb = WHITE
    tf = tb(s, 1.35, y + 0.13, 6.05, 0.85)
    para(tf, t, 13.5, True, NAVY, first=True)
    para(tf, d, 10.5, False, GRAY, before=2)
    y += 1.23

pic(s, "调研.png", 7.95, 1.45, w=4.85, h=3.64)
tf = tb(s, 7.95, 5.14, 4.85, 0.3)
para(tf, "▲ 投促局实地调研（调研记录）", 9.5, False, GRAY, align=PP_ALIGN.CENTER, first=True)

bar = card(s, 0.5, 5.55, 12.33, 1.62, NAVY, radius=0.05)
tf = tb(s, 0.9, 5.78, 11.5, 1.2)
rich(tf, [("应对方案：", 14, True, GOLD),
          ("建设区级独立平台，与市级形成「报送 + 管理」互补格局", 15, True, WHITE)], first=True)
rich(tf, [("市级系统 = 上报通道　｜　区级平台 = 全量项目底座", 12, False, WHITE)], before=8)
rich(tf, [("市级不提供接口，采用", 12, False, WHITE),
          ("市级导出数据 Excel 批量导入", 12, True, GOLD),
          ("衔接，导入时自动判重", 12, False, WHITE)], before=5)
page_no(s, 1)

# ============ P2 做什么 ============
s = slide()
title(s, "做什么：需求全景与总体设计", "15 项调研需求汇聚为四大类，落成一个平台")

tf = tb(s, 0.5, 1.25, 3.6, 0.35)
para(tf, "需求全景 · 四大类", 13, True, GOLD, first=True)

reqs = [
    ("独立部署与数据导入", "区级独立系统；市级数据 Excel 批量导入"),
    ("项目全生命周期管理", "谋划 → 在谈 → 签约 → 落地 + 退库旁路"),
    ("多主体考核", "目标设定 · 动作填报 · 审核管理 · 统计看板"),
    ("智能体应用", "自动判重 · AI 统计报告 · AI 分析总结 · 站内信"),
]
y = 1.68
for i, (t, d) in enumerate(reqs, 1):
    card(s, 0.5, y, 3.65, 1.12, BLUE_BG)
    tf = tb(s, 0.72, y + 0.14, 3.25, 0.9)
    rich(tf, [(f"0{i}  ", 13, True, GOLD), (t, 13, True, NAVY)], first=True)
    para(tf, d, 10, False, GRAY, before=3)
    y += 1.27

pic(s, "架构图.png", 4.5, 1.42, w=5.85, h=4.88)
tf = tb(s, 4.5, 6.34, 5.85, 0.3)
para(tf, "▲ 平台总体架构", 9.5, False, GRAY, align=PP_ALIGN.CENTER, first=True)

tf = tb(s, 10.6, 1.3, 2.35, 0.35)
para(tf, "设计原则", 13, True, GOLD, first=True)
prins = ["正向流转\n不回退", "一次录入 全程继承\n自动判重", "考核数据\n从业务自动生成"]
y = 1.72
for t in prins:
    c = card(s, 10.6, y, 2.35, 1.4, LIGHT, line=BORDER)
    tf = tb(s, 10.72, y + 0.18, 2.11, 1.05, MSO_ANCHOR.MIDDLE)
    lines = t.split("\n")
    para(tf, lines[0], 11.5, True, NAVY, align=PP_ALIGN.CENTER, first=True)
    if len(lines) > 1:
        para(tf, lines[1], 11.5, True, NAVY, align=PP_ALIGN.CENTER, before=2)
    y += 1.55
page_no(s, 2)

# ============ P3 场景一 ============
s = slide()
title(s, "核心场景① 项目全流程管理：从线索到落地，全程数字化留痕")

card(s, 0.5, 1.28, 12.33, 0.72, BLUE_BG, radius=0.1)
tf = tb(s, 0.8, 1.42, 11.8, 0.5)
rich(tf, [("四个阶段", 12, True, NAVY), (" 正向流转、数据自动继承、超期自动预警　•　", 12, False, DARK),
          ("在谈阶段", 12, True, NAVY), (" 决策节点管理 · 协作分派 · 项目移交 · 进展汇报　•　", 12, False, DARK),
          ("全程留痕", 12, True, NAVY), (" 进展时间线可追溯", 12, False, DARK)],
     first=True)

pic(s, "核心场景一配图.png", 1.87, 2.32, w=9.6, h=4.44)
tf = tb(s, 1.87, 6.86, 9.6, 0.3)
para(tf, "▲ 平台演示：在谈项目全流程管理", 9.5, False, GRAY, align=PP_ALIGN.CENTER, first=True)
page_no(s, 3)

# ============ P4 场景二 ============
s = slide()
title(s, "核心场景② 绩效考核闭环：园区、国企、部门差异化考核，全程数字化")

steps = [
    ("目标设定", "指标模板绑定\n批量下发、版本可回滚"),
    ("台账填报", "见商/出差/活动一事一记\n佐证材料附件上传"),
    ("审核累计", "初审（业务处室）→ 终审（局办公室）\n通过自动计入指标"),
    ("双看板", "进度看板 · 总览看板\n排名对比、动态预警"),
]
x = 0.55
for i, (t, d) in enumerate(steps):
    c = card(s, x, 1.3, 2.78, 1.05, NAVY if i == 3 else BLUE_BG, radius=0.09)
    tf = tb(s, x + 0.16, 1.4, 2.46, 0.88)
    para(tf, t, 13, True, WHITE if i == 3 else NAVY, first=True)
    lines = d.split("\n")
    for j, ln in enumerate(lines):
        para(tf, ln, 9, False, RGBColor(0xDD, 0xE5, 0xF2) if i == 3 else GRAY, before=1)
    if i < 3:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.82),
                                Inches(1.68), Inches(0.32), Inches(0.28))
        ar.fill.solid(); ar.fill.fore_color.rgb = GOLD; ar.line.fill.background()
        ar.shadow.inherit = False
    x += 3.14

tf = tb(s, 0.55, 2.52, 12.2, 0.4)
rich(tf, [("差异化考核：", 12, True, GOLD),
          ("园区 考招商成果与活动　｜　国企 考出资与产业招引　｜　部门 考落地服务与协同", 12, False, DARK)],
     first=True)

pic(s, "核心场景二配图.png", 2.17, 3.0, w=9.0, h=4.22)
tf = tb(s, 2.17, 7.22, 9.0, 0.26)
para(tf, "▲ 平台演示：绩效考核进度看板", 9.5, False, GRAY, align=PP_ALIGN.CENTER, first=True)
page_no(s, 4)

# ============ P5 场景三 AI ============
s = slide()
title(s, "核心场景③ AI 智能体：判重、报告、解析、提醒，全流程减负")

cells = [
    ("AI 自动判重", "新增/编辑/导入全场景四级判重（L1~L4），疑似重复进研判池统一裁决", "p5_duplicate.png", 2.00),
    ("智能统计 + AI 统计报告", "签约、落地、到位资金等指标实时计算；AI 将项目统计与考核统计合并生成综合报告", "p5_stats.png", 2.18),
    ("AI 分析总结", "考核佐证材料（会议纪要、签到表等）AI 自动解析文档内容、自动核算", "p5_ai.png", 1.55),
    ("站内信提醒", "任务分派、反馈催办、超期预警、审核结果自动触达，点击直达处理页", "p5_msg.png", 2.16),
]
pos = [(0.45, 1.3), (6.85, 1.3), (0.45, 4.32), (6.85, 4.32)]
CW, CH = 6.03, 2.86
for (t, d, img, ratio), (cx, cy) in zip(cells, pos):
    card(s, cx, cy, CW, CH, LIGHT, line=BORDER)
    tf = tb(s, cx + 0.2, cy + 0.12, CW - 0.4, 0.75)
    para(tf, t, 13.5, True, NAVY, first=True)
    para(tf, d, 9.5, False, GRAY, before=2)
    iw = min(3.9, 1.62 * ratio)
    ih = iw / ratio
    if ih > 1.85:
        ih = 1.85
        iw = ih * ratio
    pic(s, img, cx + (CW - iw) / 2, cy + 0.95, w=iw, h=ih)
page_no(s, 5)

# ============ P6 阶段目标与预期效果 ============
s = slide()
title(s, "阶段成果与预期效果", "与调研痛点一一呼应，效果可量化、可检验")

ones = [
    ("一份需求文档", "PRD 覆盖 4 大板块\n20+ 章节、105 项功能点"),
    ("一套演示原型", "DEMO 覆盖四大模块\n核心业务场景"),
    ("一套验收标准", "三项可量化的\n核心验收指标"),
]
x = 0.5
for t, d in ones:
    card(s, x, 1.3, 3.94, 1.22, NAVY, radius=0.08)
    tf = tb(s, x + 0.25, 1.44, 3.44, 0.95)
    para(tf, t, 14, True, GOLD, first=True)
    lines = d.split("\n")
    for ln in lines:
        para(tf, ln, 10.5, False, WHITE, before=1)
    x += 4.2

tf = tb(s, 0.5, 2.78, 6.0, 0.35)
para(tf, "调研痛点 → 预期效果", 13, True, GOLD, first=True)

maps = [
    ("进度靠电话催、表格记，看不到全量",
     "全流程可追溯管理", "谋划到落地全程留痕，项目信息全面、准确，进度一屏可见"),
    ("考核线下人工统计，耗时耗力",
     "考核全线上化", "填报、审核、统计全程线上，节约时间成本，佐证材料线上留存"),
    ("多头填报、重复与口径不一",
     "数据一次录入、全程继承", "AI 判重拦截重复项目，统计口径统一、可验算"),
]
y = 3.18
for pain, eff, d in maps:
    card(s, 0.5, y, 4.6, 0.86, GRAY_BG)
    tf = tb(s, 0.7, y + 0.1, 4.25, 0.68, MSO_ANCHOR.MIDDLE)
    para(tf, pain, 10.5, False, GRAY, first=True)
    ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.22), Inches(y + 0.28),
                            Inches(0.5), Inches(0.3))
    ar.fill.solid(); ar.fill.fore_color.rgb = GOLD; ar.line.fill.background()
    ar.shadow.inherit = False
    card(s, 5.85, y, 6.98, 0.86, BLUE_BG)
    tf = tb(s, 6.08, y + 0.09, 6.55, 0.7)
    rich(tf, [(eff, 12, True, NAVY), ("　" + d, 10.5, False, GRAY)], first=True)
    y += 0.98

kpis = [("≥90%", "AI 判重准确率"), ("≥90%", "AI 数据引用准确率"), ("≥90%", "统计口径人工验算一致率")]
x = 0.5
for v, t in kpis:
    card(s, x, 6.28, 3.94, 0.92, WHITE, line=GOLD, radius=0.1)
    tf = tb(s, x + 0.25, 6.42, 3.44, 0.65, MSO_ANCHOR.MIDDLE)
    rich(tf, [(v + "  ", 20, True, GOLD), (t, 12.5, True, NAVY)], first=True)
    x += 4.2
page_no(s, 6)

prs.save(OUT)
print("saved:", OUT)
